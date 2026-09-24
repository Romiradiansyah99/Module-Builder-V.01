"""
WORD INJECTOR - Rakit draft_json menjadi file .docx fisik
=========================================================
Contoh penggunaan library docxtpl:

    from docxtpl import DocxTemplate

    tpl = DocxTemplate("template_kemnaker.docx")  # template berisi {{ judul_modul }} dll.
    tpl.render({"judul_modul": "Instalasi PLTS", "konten": "..."})
    tpl.save("output/hasil.docx")

Modul ini memakai template resmi Kemnaker (database/template_kemnaker.docx
- dibangun dari Templet Modul_2024 resmi, lihat tools/template_builder.py).
Tag yang tidak diisi LLM diberi nilai default "--- (perlu dilengkapi
reviewer) ---" agar render tidak error dan reviewer bisa melihat bagian
mana yang bolong.

Ronde 11 (kesesuaian template + gambar wajib):
  * pengetahuan_content dirakit via docxtpl Subdoc -> paragraf Word ASLI:
    baris bergaya judul subbab ("1. ..." / "1.1 ...") dibuat BOLD, blok
    ```mermaid ... ``` dirender jadi GAMBAR PNG (tools/mermaid_render.py)
    dan disisipkan sebagai gambar - sesuai tuntutan user bahwa flowchart
    tidak boleh hilang dari .docx.
  * lik_gambar_kerja yang berisi blok mermaid juga dirender jadi gambar
    (docxtpl InlineImage); jika bukan skrip mermaid, dimasukkan sebagai teks.
  * key row-loop docxtpl (elemen_rows, bahan_rows, cek_observasi_rows,
    cek_hasil_rows, kamus_rows, referensi_rows) dinormalkan: baris yang
    kurang key dilengkapi string kosong agar render tak pernah Undefined.
"""

import os
import re
from pathlib import Path
from typing import List, Optional

from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt
from docxtpl import DocxTemplate, InlineImage, Listing
from docxtpl.subdoc import Subdoc

from agents.state import GlobalState, ModuleState
from tools.doc_utils import (
    OUTPUT_DIR,
    TEMPLATE_PATH,
    cap_first,
    get_template_tags,
    pengetahuan_items,
)
from tools.mermaid_render import render_mermaid

# Ronde 13: key draft BUKAN tag template - query pencarian gambar internet
# yang ditulis Agent 2 (dieksekusi tools/image_search.py saat injeksi).
IMAGE_QUERY_KEYS = ("cover_image_query", "image_queries")

# Nilai default untuk tag template yang tidak terisi LLM
DEFAULT_MISSING = "--- (perlu dilengkapi reviewer) ---"

# Key row-loop docxtpl ({%tr %}) - tidak terdeteksi regex tag sederhana
ROW_KEYS = (
    "elemen_rows",
    "bahan_rows",
    "cek_observasi_rows",
    "cek_hasil_rows",
    "kamus_rows",
    "referensi_rows",
)

# Key wajib tiap baris loop - dilengkapi "" agar render tidak Undefined
_ROW_REQUIRED = {
    "elemen_rows": ("elemen_no", "elemen", "kuk_no", "kuk", "indikator",
                    "pengetahuan", "keterampilan", "durasi"),
    "bahan_rows": ("no", "nama", "spek", "jumlah"),
    "cek_observasi_rows": ("langkah", "acuan"),
    "cek_hasil_rows": ("no", "aspek", "standar"),
    "kamus_rows": ("no", "istilah", "arti"),
    "referensi_rows": ("no", "url"),
}

_HEADING_RE = re.compile(r"^\d+(\.\d+)*\.?\s+\S")  # "1. Judul" / "1.1 Judul"
_MERMAID_RE = re.compile(r"```mermaid\s*(.*?)```", re.DOTALL)


def _slugify(text: str) -> str:
    """Judul modul -> nama file yang aman (tanpa karakter ilegal Windows)."""
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")[:60] or "modul"


def _strip_heading(heading: str) -> str:
    """Hapus nomor di awal judul subbab: "1. Merakit Panel Surya" ->
    "Merakit Panel Surya". Dipakai ronde 16 untuk judul caption & fallback
    query gambar bila Agent 2 tidak menyediakan entri foto utk subbab itu."""
    return re.sub(r"^\d+(\.\d+)*\.?\s+", "", heading or "").strip()


def _render_mermaid_safe(code: str, out_path: Path) -> Optional[Path]:
    """Render skrip mermaid -> PNG; None kalau gagal (jangan pernah crash
    injection karena skrip mermaid yang tidak sempurna)."""
    try:
        return render_mermaid(code, out_path)
    except Exception as exc:  # noqa: BLE001
        print(f"[word_injector] Mermaid gagal dirender ({exc}) - lewati gambar")
        return None


def _fit_png(png: Optional[Path], max_w_mm: float,
             max_h_mm: float) -> Optional[float]:
    """Lebar tampilan (mm) gambar PNG agar PAS KOTAK max_w x max_h mm
    (ronde 12, komentar reviewer: "terlalu besar, tolong dikecilkan dan
    disesuaikan"). Ukuran piksel dibaca via PIL; rasio aspek dipertahankan
    dan gambar tidak pernah melebihi kedua batas. None bila PNG tak terbaca."""
    if png is None:
        return None
    try:
        from PIL import Image

        with Image.open(str(png)) as im:
            px_w, px_h = im.size
    except Exception:  # noqa: BLE001 - PIL tak terbaca: pakai lebar aman
        return min(max_w_mm, 120.0)
    if px_w <= 0 or px_h <= 0:
        return min(max_w_mm, 120.0)
    aspect = px_w / px_h
    return round(min(max_w_mm, max_h_mm * aspect), 1)


# Batas ukuran gambar di docx final (mm) - kertas A4 margin template:
# area teks ~150mm. Contoh modul memakai gambar 43-129mm; gambar raksasa
# 145x350mm/150x594mm pada output sebelumnya dikritik reviewer.
_IMG_PENGETAHUAN_MAX = (135.0, 100.0)  # diagram flowchart pengetahuan
_IMG_GAMBAR_KERJA_MAX = (120.0, 100.0)  # gambar kerja LIK
_IMG_COVER_MAX = (120.0, 130.0)  # foto cover page (ronde 13)
_COVER_Y_MM = 72.0  # jarak dari atas kertas: judul berakhir ~47mm, footer 268mm


def _fetch_image_safe(query: str, out_path: Path, mode: str = "cari") -> Optional[Path]:
    """Gambar utk query dengan mode pilihan Agent 2 (ronde 13b):
    - "generate": AI Replicate (google/nano-banana-v2) dulu -> gagal: cari internet.
    - "cari"    : foto nyata internet dulu -> gagal: fallback AI Replicate
      (agar tiap subbab/cover dijamin punya gambar; Replicate dijalankan
      hanya bila pencarian kosong).
    Selalu return path atau None - kegagalan internet/API TIDAK boleh
    menggagalkan injeksi Word."""
    query = str(query or "").strip()
    if not query:
        return None

    def _generate() -> Optional[Path]:
        try:
            from tools.image_gen import generate_image

            png = generate_image(query, out_path)
            if png:
                return Path(png)
        except Exception as exc:  # noqa: BLE001
            print(f"[word_injector] Generator Replicate gagal ({exc}) - lewati")
        return None

    def _search() -> Optional[Path]:
        try:
            from tools.image_search import fetch_image

            got = fetch_image(query, out_path)
            if got:
                print(f"[word_injector] Gambar internet OK ({got.get('source')}): "
                      f"{query[:60]}")
                return Path(got["path"])
        except Exception as exc:  # noqa: BLE001
            print(f"[word_injector] Pencari gambar gagal ({exc}) - lewati")
        return None

    if mode == "generate":
        return _generate() or _search()
    return _search() or _generate()


def _normalize_image_queries(value) -> dict:
    """draft_json["image_queries"] -> {"<no subbab>": {"query","judul"}}.
    Bentuk mentah dari LLM: list of {"subbab": "1", "query": "...",
    "judul": "..."}. Entri rusak dibuang - sisanya tetap dipakai."""
    out: dict = {}
    if not isinstance(value, list):
        return out
    for item in value:
        if not isinstance(item, dict):
            continue
        query = str(item.get("query") or "").strip()
        if not query:
            continue
        num = re.sub(r"\.0$", "", str(item.get("subbab") or item.get("no") or "")).strip().rstrip(".")
        if not num:
            continue
        out[num] = {
            "query": query,
            "judul": str(item.get("judul") or query).strip(),
            # ronde 13b: "cari" (foto internet, default) | "generate" (AI Replicate)
            "mode": str(item.get("mode") or "cari").strip() or "cari",
        }
    return out


def _inline_to_anchor(inline, x_emu: int, y_emu: int, docpr_id: int):
    """Ubah wp:inline (gambar inline) menjadi wp:anchor MELAYANG dengan
    posisi absolut dari tepi kertas (ronde 13, foto cover page). Semua anak
    extent/docPr/graphic dipindahkan utuh - r:id hub tetap valid."""
    from lxml import etree

    from docx.oxml import parse_xml

    def _xml(el):
        return etree.tostring(el, encoding="unicode")

    extent = inline.find(qn("wp:extent"))
    docpr = inline.find(qn("wp:docPr"))
    frame = inline.find(qn("wp:cNvGraphicFramePr"))
    graphic = inline.find(qn("a:graphic"))
    docpr.set("id", str(docpr_id))
    docpr.set("name", f"Cover Image {docpr_id}")
    return parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="114300" distR="114300" simplePos="0" '
        'relativeHeight="251658240" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{x_emu}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{y_emu}</wp:posOffset></wp:positionV>'
        + _xml(extent)
        + '<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        + _xml(docpr)
        + (_xml(frame) if frame is not None else "")
        + _xml(graphic)
        + "</wp:anchor>"
    )


def _add_cover_image(doc, img_path: Path) -> None:
    """Sisipkan foto ke COVER PAGE (ronde 13, komentar reviewer: cover harus
    ada gambar relevan dengan judul). Sebagai ANCHOR melayang posisi absolut
    (relativeFrom=page) - tidak menggeser aliran teks cover yang dibangun
    textbox melayang (judul di y=4-47mm, footer y=268mm)."""
    from PIL import Image

    with Image.open(str(img_path)) as im:
        aspect = im.width / max(im.height, 1)
    max_w, max_h = _IMG_COVER_MAX
    w_mm = min(max_w, max_h * aspect)
    h_mm = w_mm / max(aspect, 0.01)
    para = doc.paragraphs[0]  # paragraf pertama = section cover
    run = para.add_run()
    run.add_picture(str(img_path), width=Mm(w_mm), height=Mm(h_mm))
    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    ids = [int(el.get("id")) for el in doc.element.body.iter(qn("wp:docPr"))]
    new_id = (max(ids) + 1) if ids else 901
    x_emu = int(((210.0 - w_mm) / 2.0) * 36000)  # tengah kertas A4
    y_emu = int(_COVER_Y_MM * 36000)
    drawing.replace(inline, _inline_to_anchor(inline, x_emu, y_emu, new_id))
    print(f"[word_injector] Foto cover disisipkan ({w_mm:.0f}x{h_mm:.0f}mm)")


def _canonical_subbab(syllabus_rows) -> tuple:
    """Judul subbab/sub-subbab LEGAL dari silabus modul (set ternormalisasi).
    Ronde 12: HANYA baris yang persis redaksi elemen/butir pengetahuan
    silabus boleh jadi heading - baris list bernomor di tengah paragraf
    ("1. Sambut truk...") tidak boleh masuk daftar isi Word."""
    subs, subsubs = set(), set()
    for r in syllabus_rows or []:
        if not isinstance(r, dict):
            continue
        try:
            n = int(str(r.get("elemen_no", "")).strip().rstrip("."))
        except ValueError:
            continue
        subs.add(re.sub(r"\s+", " ", f"{n}. {r.get('elemen', '')}".strip().lower()).strip(" .;"))
        for i, item in enumerate(
                pengetahuan_items(str(r.get("pengetahuan") or "")), start=1):
            item_clean = cap_first(item.strip())
            subsubs.add(re.sub(r"\s+", " ", f"{n}.{i} {item_clean}".lower()).strip(" .;"))
    return subs, subsubs


def _build_pengetahuan_subdoc(tpl: DocxTemplate, text: str, tmp_dir: Path,
                              syllabus_rows=None, photo_queries=None):
    """Rakit nilai pengetahuan_content menjadi Subdoc berisi paragraf asli:
    judul subbab -> style Heading 2 (masuk daftar isi Word); blok
    ```mermaid``` -> gambar PNG terukur + CAPTION "Gambar N. <judul subbab>"
    (ronde 12, komentar reviewer: gambar/flowchart wajib berjudul).

    Ronde 13: photo_queries = {"<no subbab>": {"query","judul"}} dari
    Agent 2 - SETELAH tiap judul subbab disisipkan FOTO nyata dari internet
    (query dieksekusi tools/image_search.py) dengan caption yang memakai
    SATU urutan "Gambar N." bersama diagram mermaid."""
    sub = Subdoc(tpl)
    canon_subs, canon_subsubs = _canonical_subbab(syllabus_rows)
    tmp_dir.mkdir(parents=True, exist_ok=True)
    pos = 0
    img_i = 0  # SATU urutan "Gambar N." utk diagram + foto (ronde 13)

    def _add_caption(judul: str) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        cap = sub.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap.add_run(f"Gambar {img_i}. {judul}")
        cap_run.font.name = "Bookman Old Style"
        cap_run.font.size = Pt(11)

    def _add_picture_centered(png, max_box, caption=None) -> bool:
        """Gambar pada PARAGRAF SENDIRI, rata-tengah (ronde 16: posisi & lebar
        konsisten, bukan inline di paragraf heading yang rata-kiri), dengan
        caption "Gambar N." opsional. Return True bila gambar berhasil
        dipasang; pemanggil mengelola counter `img_i`."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        width_mm = _fit_png(png, *max_box)
        if not width_mm:
            return False
        p = sub.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(png), width=Mm(width_mm))
        if caption:
            _add_caption(caption)
        return True

    def _maybe_photo(heading: str) -> None:
        """Foto/ilustrasi ter-center tepat di bawah judul subbab (komentar
        reviewer ronde 13: "masih ga ada gambar" - diagram saja tidak cukup;
        ronde 16: tiap subbab DIJAMIN punya gambar, rata-tengah, ukuran
        konsisten). Counter "Gambar N." hanya naik saat foto BENAR-BENAR
        terpasang - kalau tidak, caption mermaid mulai dari "Gambar 2." tanpa
        "Gambar 1." (bug terukur pada smoke ronde 13)."""
        nonlocal img_i
        m = re.match(r"^(\d+)\.\s", heading or "")
        if not m:
            return
        subbab_no = m.group(1)
        spec = (photo_queries or {}).get(subbab_no)
        # Ronde 16 ketersediaan: subbab tanpa entri query tetap wajib punya
        # gambar -> fallback query dari judul subbab itu sendiri.
        query = (spec or {}).get("query") or _strip_heading(heading)
        if not query:
            return
        png = _fetch_image_safe(
            query, tmp_dir / f"foto_{subbab_no}.png",
            mode=str((spec or {}).get("mode") or "cari"),
        )
        if png is not None:
            img_i += 1
            judul = ((spec or {}).get("judul") or (spec or {}).get("query")
                     or _strip_heading(heading) or "Gambar kerja")
            _add_picture_centered(png, _IMG_PENGETAHUAN_MAX, judul)

    def _add_lines(chunk: str) -> None:
        nonlocal last_heading
        for line in chunk.split("\n"):
            new_heading = _subdoc_add_line(sub, line, last_heading,
                                           canon_subs, canon_subsubs)
            if new_heading != last_heading and new_heading:
                last_heading = new_heading
                _maybe_photo(last_heading)

    last_heading = ""
    for m in _MERMAID_RE.finditer(text):
        # teks sebelum blok mermaid
        _add_lines(text[pos : m.start()])
        code = m.group(1)
        img_i += 1   # reserve nomor "Gambar N." walau render gagal berhasil
        png = _render_mermaid_safe(code, tmp_dir / f"mermaid_{img_i}.png")
        if png is not None:
            # Ronde 16: diagram tetap di posisi inline tempat penulisan (di
            # dalam urutan teks subbab), tetapi di-center + caption + lebar
            # konsisten (sebelumnya rata-kiri inline di paragraf teks).
            judul = _strip_heading(last_heading) or "Diagram alur kerja"
            if not _add_picture_centered(png, _IMG_PENGETAHUAN_MAX, judul):
                # PIL tak bisa membaca dimensi -> tetap center, lebar tetap.
                from docx.enum.text import WD_ALIGN_PARAGRAPH

                p = sub.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(png), width=Inches(5.7))
                _add_caption(judul)
        else:
            # renderer gagal -> simpan skripnya sebagai teks agar tidak hilang
            for line in ("```mermaid" + code + "```").split("\n"):
                _subdoc_add_line(sub, line, last_heading,
                                 canon_subs, canon_subsubs)
        pos = m.end()
    _add_lines(text[pos:])
    return sub


def _subdoc_add_line(sub, line: str, last_heading: str = "",
                     canon_subs=frozenset(), canon_subsubs=frozenset()) -> str:
    """Satu baris -> satu paragraf Subdoc. Return judul subbab terakhir
    terbaru (untuk caption gambar).

    - Baris "N. <redaksi elemen silabus PERSIS>" -> style Heading 2 + run
      Bookman 12 bold (masuk daftar isi TOC \\o "1-2"). Baris bernomor lain
      (list di tengah paragraf) tetap teks biasa - jangan memenuhi TOC.
    - Baris "N.M <redaksi pengetahuan PERSIS>" -> paragraf BOLD (sub-subbab,
      tidak memenuhi TOC).
    - Baris isi -> paragraf justify Bookman 12.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    line = line.rstrip()
    if not line.strip():
        return last_heading  # paragraf kosong tidak diperlukan
    stripped = line.strip()
    m = re.match(r"^(\d+(?:\.\d+)*)\.?\s+(.*\S)$", stripped)
    if m:
        # titik pemisah habis dimakan regex pada "1. judul" - pulihkan agar
        # key cocok dengan canonical "1. judul" (subbab) dan "1.1 judul"
        num = m.group(1) if "." in m.group(1) else m.group(1) + "."
        key = re.sub(r"\s+", " ", f"{num} {m.group(2)}".strip().lower())
        if "." in m.group(1):
            if key in canon_subsubs:
                p = sub.add_paragraph()
                run = p.add_run(stripped)
                run.bold = True
                run.font.name = "Bookman Old Style"
                run.font.size = Pt(12)
                return stripped
        elif key in canon_subs:
            p = sub.add_paragraph(style="Heading 2")
            run = p.add_run(stripped)
            run.bold = True
            run.font.name = "Bookman Old Style"
            run.font.size = Pt(12)
            return stripped
    p = sub.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(stripped)
    run.font.name = "Bookman Old Style"
    run.font.size = Pt(12)
    return last_heading


def _normalize_rows(key: str, value) -> list:
    """Pastikan row-loop selalu list of dict dengan key lengkap."""
    if not isinstance(value, list):
        return []
    required = _ROW_REQUIRED.get(key, ())
    rows = []
    for row in value:
        if isinstance(row, dict):
            clean = {k: ("" if v is None else str(v)) for k, v in row.items()}
            for k in required:
                clean.setdefault(k, "")
            rows.append(clean)
    return rows


def inject_module(module: ModuleState, output_dir: Path = None) -> Path:
    """Render SATU modul menjadi file .docx. Return path hasil.

    Raises:
        RuntimeError: jika template tidak ditemukan / draft kosong.
    """
    if not TEMPLATE_PATH.exists():
        raise RuntimeError(f"Template tidak ditemukan: {TEMPLATE_PATH}")

    draft_json = module.get("draft_json") or {}
    if not draft_json:
        raise RuntimeError(f'Modul {module.get("module_id", "?")} tidak punya draft_json.')

    out_dir = Path(output_dir) if output_dir else OUTPUT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)

    # --- Injeksi dictionary ke template ---
    tpl = DocxTemplate(str(TEMPLATE_PATH))

    # Pre-fill SEMUA tag template dengan nilai default -> render tidak error,
    # dan reviewer bisa melihat bagian mana yang belum terisi.
    context = {tag: DEFAULT_MISSING for tag in get_template_tags()}

    # Timpa dengan hasil Agent 2 (hanya nilai non-kosong yang menimpa default).
    # String multi-baris dibungkus docxtpl.Listing agar newline menjadi paragraf
    # asli Word (bukan karakter newline mentah yang diruntuhkan).
    for key, value in draft_json.items():
        if key in ROW_KEYS:
            rows = _normalize_rows(key, value)
            context[key] = rows  # loop kosong -> tabel hanya header (aman)
        elif isinstance(value, str) and value.strip():
            context[key] = Listing(value) if "\n" in value.strip() else value
        elif value:  # nilai non-string lain: lempar mentah
            context[key] = value

    # Ronde 11: pengetahuan_content & lik_gambar_kerja perlakuan khusus
    # (paragraf bold + gambar mermaid) agar .docx final punya DIAGRAM.
    tmp_dir = out_dir / "_mermaid"
    content = str(draft_json.get("pengetahuan_content") or "").strip()
    # Judul subbab legal diambil dari silabus state; bila kosong (modul lama/
    # demo) - turunkan dari elemen_rows draf (strip nomor pada kolom elemen).
    syllabus_rows = module.get("syllabus_rows") or [
        {**r, "elemen": re.sub(r"^\d+\.\s*", "", str(r.get("elemen", "")))}
        for r in (draft_json.get("elemen_rows") or []) if isinstance(r, dict)
    ]
    if content:
        try:
            context["pengetahuan_content"] = _build_pengetahuan_subdoc(
                tpl, content, tmp_dir, syllabus_rows=syllabus_rows,
                photo_queries=_normalize_image_queries(
                    draft_json.get("image_queries")
                ),
            )
        except Exception as exc:  # noqa: BLE001 - fallback teks, jangan gagal
            print(f"[word_injector] Subdoc pengetahuan gagal ({exc}) - pakai Listing")
            context["pengetahuan_content"] = Listing(content)

    gambar = str(draft_json.get("lik_gambar_kerja") or "").strip()
    if gambar and "```" in gambar:
        # Ronde 12: gambar kerja di-scan blok mermaid pertama lalu diskalakan
        # pas kotak 120x100mm (dulu Mm(150) penuh -> 150x594mm, dikritik
        # reviewer "terlalu besar, tolong dikecilkan dan disesuaikan").
        m = _MERMAID_RE.search(gambar)
        png = _render_mermaid_safe(
            m.group(1) if m else gambar, tmp_dir / "gambar_kerja.png"
        )
        if png is not None:
            width_mm = _fit_png(png, *_IMG_GAMBAR_KERJA_MAX) or 120.0
            context["lik_gambar_kerja"] = InlineImage(
                tpl, str(png), width=Mm(width_mm)
            )
        else:
            context["lik_gambar_kerja"] = Listing(gambar)
    elif gambar:
        context["lik_gambar_kerja"] = Listing(gambar) if "\n" in gambar else gambar

    # Selalu pastikan identitas modul terisi
    context["module_title"] = module.get("module_title", context.get("module_title", ""))
    if "module_id" in context:
        context["module_id"] = module.get("module_id", "")

    tpl.render(context)

    # Ronde 13: foto cover page - diambil dari internet, relevan dengan
    # judul modul (query ditulis Agent 2; fallback = judul modul).
    # Ronde 13b: cover_image_query boleh STRING atau {"query","mode"}.
    try:
        cv = draft_json.get("cover_image_query")
        if isinstance(cv, dict):
            cover_query = str(cv.get("query") or module.get("module_title") or "").strip()
            cover_mode = str(cv.get("mode") or "cari").strip() or "cari"
        else:
            cover_query = str(cv or module.get("module_title") or "").strip()
            cover_mode = "cari"
        cover_png = _fetch_image_safe(cover_query, tmp_dir / "cover.png",
                                      mode=cover_mode)
        if cover_png is not None:
            # PENTING: tpl.docx (bukan get_docx()) - get_docx() me-reload
            # dokumen dari template mentah saat is_rendered=True (seluruh
            # hasil render hilang, docx final berisi tag {{ }} tak terisi).
            _add_cover_image(tpl.docx, cover_png)
    except Exception as exc:  # noqa: BLE001 - cover polos tetap jalan
        print(f"[word_injector] Foto cover gagal ({exc}) - lewati")

    # --- Simpan hasil (atomik) ---
    filename = f'{module.get("module_id", "M?")}_{_slugify(module.get("module_title", "modul"))}.docx'
    out_path = out_dir / filename
    # Tulis ke file .tmp lalu os.replace -> render paralel antar user tidak
    # pernah meninggalkan .docx setengah jadi yang bisa terunduh korup.
    # Windows (dev): file target masih terbuka di Word -> Access denied;
    # fallback simpan langsung agar demo/tes lokal tetap jalan.
    tmp_path = out_path.with_suffix(".docx.tmp")
    tpl.save(str(tmp_path))
    try:
        os.replace(str(tmp_path), str(out_path))
    except PermissionError:
        os.unlink(str(tmp_path))
        tpl.save(str(out_path))
        print(f"[word_injector] Replace atomik terblokir (file terbuka?) - simpan langsung: {out_path}")
    print(f"[word_injector] Saved: {out_path}")
    return out_path


def inject_all_modules(state: GlobalState, output_dir: Path = None) -> List[Path]:
    """Loop semua modul yang sudah PASS lalu rakit .docx satu per satu.

    Modul yang gagal render tidak menghentikan modul lain (error dicatat,
    proses lanjut) supaya satu modul bermasalah tidak membatalkan batch.
    """
    results: List[Path] = []
    for module in state.get("modules") or []:
        try:
            results.append(inject_module(module, output_dir))
        except Exception as exc:  # noqa: BLE001 - jangan gagalkan batch
            print(f"[word_injector] GAGAL modul {module.get('module_id')}: {exc}")
    return results


# ----------------------------------------------------------------------
# Demo/contoh penggunaan: python tools/word_injector.py
# ----------------------------------------------------------------------
if __name__ == "__main__":
    demo_pengetahuan = (
        "1. Persiapan\n"
        "Perencanaan dilakukan untuk memastikan seluruh kebutuhan tersedia "
        "sebelum pekerjaan dimulai. Instruktur menjelaskan urutan kerja.\n"
        "1.1 Jenis alat\n"
        "Perencanaan adalah proses menentukan tujuan dan langkah kerja.\n"
        "```mermaid\n"
        "flowchart TD\n"
        "A[Mulai] --> B[Menyiapkan alat]\n"
        "B --> C[Cek kelengkapan]\n"
        "C --> D[Melaksanakan pekerjaan]\n"
        "D --> E[Membuat laporan]\n"
        "```\n"
        "2. langkah list bernomor tidak jadi heading\n"
        "Sarana dan prasarana dicek kelengkapannya sebelum digunakan."
    )
    demo_module = ModuleState(
        module_id="DEMO",
        module_title="Mengoperasikan Kondenser",
        syllabus_content="| Elemen Kompetensi | KUK | Alokasi Waktu |\n|---|---|---|\n| Persiapan | Alat lengkap | 1 JP |",
        syllabus_rows=[
            {"elemen_no": "1.", "elemen": "Persiapan", "kuk_no": "1.1",
             "kuk": "Alat lengkap", "indikator": "Alat teridentifikasi",
             "pengetahuan": "Penjelasan tentang:\n1. Jenis alat\n2. Fungsi alat",
             "keterampilan": "Mengidentifikasi alat", "durasi": "1 JP"},
        ],
        draft_json={
            "pengetahuan_content": demo_pengetahuan,
            "lik_gambar_kerja": (
                "```mermaid\nflowchart TD\nA[Mulai] --> B[Cek alat]\nB --> C[Kerjakan]\n```"
            ),
            "cover_image_query": "steam power plant condenser",
            "image_queries": [
                {"subbab": "1", "query": "industrial hand tools set",
                 "judul": "Jenis-jenis alat tukang industri"},
            ],
            "elemen_rows": [
                {"elemen_no": "1.", "elemen": "Persiapan", "kuk_no": "1.1",
                 "kuk": "Alat lengkap", "indikator": "Alat teridentifikasi",
                 "pengetahuan": "1. Jenis alat\n2. Fungsi alat",
                 "keterampilan": "Mengidentifikasi alat", "durasi": "1 JP"},
            ],
            "bahan_rows": [{"no": "1", "nama": "Kertas HVS", "spek": "A4 80gsm", "jumlah": "10 lembar"}],
            "kamus_rows": [{"no": "a.", "istilah": "Kondenser", "arti": "Komponen pendingin"}],
            "referensi_rows": [{"no": "a.", "url": "https://kbbi.co.id"}],
        },
        evaluator_feedback="",
        status_evaluasi="PASS",
        iteration_count=1,
    )
    print(f"Tag template ditemukan: {get_template_tags()}")
    path = inject_module(demo_module, output_dir=OUTPUT_DIR / "demo")
    print(f"Demo render sukses -> {path}")